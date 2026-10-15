from pathlib import Path

import pytest
from docx import Document as DocxDocument

import src.ingestion as ing

ROOT = Path(__file__).resolve().parent.parent


@pytest.fixture
def captured(monkeypatch):
    calls = []
    monkeypatch.setattr(ing, "add_documents", lambda chunks, ids, only_empty=False: calls.append((chunks, ids, only_empty)) or {"pinecone": len(chunks)})
    return calls


def test_markdown_is_loaded_with_metadata(settings, tmp_path):
    p = tmp_path / "runbook.md"
    p.write_text("# Title\n\nRestart the worker.", encoding="utf-8")
    docs = ing.load_file(p)
    assert len(docs) == 1
    assert docs[0].metadata["title"] == "runbook.md"
    assert docs[0].metadata["document_name"] == "runbook.md"


def test_docx_is_loaded(settings, tmp_path):
    p = tmp_path / "sop.docx"
    d = DocxDocument()
    d.add_paragraph("Rollback when error rate exceeds 5%.")
    d.add_paragraph("")
    d.add_paragraph("Notify the incident channel.")
    d.save(p)
    text = ing.load_file(p)[0].page_content
    assert "Rollback" in text and "incident channel" in text


def test_pdf_pages_keep_page_numbers(settings):
    pdf = ROOT / "uploads" / "novacore_company_information.pdf"
    if not pdf.exists():
        pytest.skip("sample PDF not present")
    docs = ing.load_file(pdf)
    assert len(docs) >= 2
    assert [d.metadata["page"] for d in docs] == list(range(len(docs)))
    assert any("ERR_CONN_8421" in d.page_content for d in docs)


def test_unsupported_extension(settings, tmp_path):
    p = tmp_path / "x.csv"
    p.write_text("a,b")
    with pytest.raises(ValueError, match="Unsupported"):
        ing.load_file(p)


def test_empty_file_is_rejected(settings, tmp_path, captured):
    p = tmp_path / "empty.txt"
    p.write_text("   \n ")
    with pytest.raises(ValueError, match="No readable text"):
        ing.ingest_file(p)
    assert captured == []


def test_chunking_ids_are_stable_unique_and_safe(settings, tmp_path, captured):
    p = tmp_path / "Big Runbook (v2).md"
    p.write_text("\n\n".join(f"Step {i}: " + "check the service health carefully. " * 12 for i in range(20)), encoding="utf-8")
    n = ing.ingest_file(p)
    chunks, ids, _ = captured[0]
    assert n == len(chunks) == len(ids) > 3
    assert len(set(ids)) == len(ids)
    assert all(len(c.page_content) <= 800 for c in chunks)
    assert all(i.startswith("Big-Runbook--v2--") for i in ids)
    assert [c.metadata["chunk_index"] for c in chunks] == list(range(n))
    ing.ingest_file(p)
    assert captured[1][1] == ids


def test_directory_ingestion_batches_all_supported_files(settings, tmp_path, captured):
    (tmp_path / "a.md").write_text("alpha runbook")
    (tmp_path / "b.txt").write_text("beta sop")
    (tmp_path / "c.png").write_bytes(b"\x89PNG")
    assert [p.name for p in ing.list_documents(tmp_path)] == ["a.md", "b.txt"]
    assert ing.ingest_directory(tmp_path, only_empty=True) == {"pinecone": 2}
    assert captured[0][2] is True
    assert ing.ingest_directory(tmp_path / "missing") == {}


def test_bundled_runbooks_are_ingestible(settings, captured):
    total = ing.ingest_directory(ROOT / "documents")
    assert total["pinecone"] >= 3
