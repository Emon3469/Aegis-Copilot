from hashlib import sha256
from pathlib import Path
from typing import Dict, List

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from src.config import get_settings
from src.vectorstore import add_documents

SUPPORTED = {".pdf", ".txt", ".md", ".docx"}

_splitter = RecursiveCharacterTextSplitter(
    chunk_size=800,
    chunk_overlap=160,
    separators=["\n\n", "\n", ". ", " ", ""],
)


def _load_pdf(path: Path) -> List[Document]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    docs = []
    for i, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        if text:
            docs.append(Document(page_content=text, metadata={"source": path.name, "page": i}))
    return docs


def _load_docx(path: Path) -> List[Document]:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return [Document(page_content=text, metadata={"source": path.name})] if text else []


def _load_text(path: Path) -> List[Document]:
    text = path.read_text(encoding="utf-8", errors="ignore").strip()
    return [Document(page_content=text, metadata={"source": path.name})] if text else []


def load_file(path: Path) -> List[Document]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        docs = _load_pdf(path)
    elif ext in {".txt", ".md"}:
        docs = _load_text(path)
    elif ext == ".docx":
        docs = _load_docx(path)
    else:
        raise ValueError(f"Unsupported file type: {ext}. Use PDF, TXT, MD, or DOCX.")
    for d in docs:
        d.metadata["title"] = path.name
        d.metadata["document_name"] = path.name
    return docs


def _chunk_id(path: Path, chunk: Document, position: int) -> str:
    digest = sha256(f"{path.name}|{position}|{chunk.page_content}".encode("utf-8")).hexdigest()[:24]
    stem = "".join(c if c.isalnum() or c in "-_" else "-" for c in path.stem)[:60]
    return f"{stem}-{position}-{digest}"


def _chunks(path: Path):
    docs = load_file(path)
    chunks = _splitter.split_documents(docs)
    ns = get_settings().pinecone_namespace
    for i, chunk in enumerate(chunks):
        chunk.metadata["chunk_index"] = i
        chunk.metadata["knowledge_base"] = ns
    return chunks, [_chunk_id(path, c, i) for i, c in enumerate(chunks)]


def ingest_file(path: Path) -> int:
    chunks, ids = _chunks(path)
    if not chunks:
        raise ValueError("No readable text was found in this file.")
    add_documents(chunks, ids)
    return len(chunks)


def list_documents(directory: Path) -> List[Path]:
    if not directory.exists():
        return []
    return [p for p in sorted(directory.iterdir()) if p.is_file() and p.suffix.lower() in SUPPORTED]


def ingest_directory(directory: Path, only_empty: bool = False) -> Dict[str, int]:
    chunks, ids = [], []
    for path in list_documents(directory):
        c, i = _chunks(path)
        chunks.extend(c)
        ids.extend(i)
    if not chunks:
        return {}
    return add_documents(chunks, ids, only_empty=only_empty)


def namespace() -> str:
    return get_settings().pinecone_namespace
